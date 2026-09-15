"""
Membrete institucional en los archivos exportados.

El riesgo que cubren estos tests no es que el membrete se vea feo: es que
al empujar los datos cuatro filas hacia abajo se rompa el round-trip, que
es la funcionalidad central del sistema.
"""

import os
import tempfile
from datetime import date

from django.contrib.auth import get_user_model
from django.test import TestCase, override_settings
from django.urls import reverse
from openpyxl import load_workbook

from tracker import membrete
from tracker.exportador import FILA_DATOS, FILA_ENCABEZADOS, exportar_tareas, exportar_tickets
from tracker.importer import es_archivo_de_tareas, importar, importar_tareas, leer_filas
from tracker.models import Tarea, Ticket


def a_disco(contenido, sufijo=".xlsx"):
    tmp = tempfile.NamedTemporaryFile(delete=False, suffix=sufijo)
    tmp.write(contenido)
    tmp.close()
    return tmp.name


class ConfiguracionTest(TestCase):
    def test_los_datos_salen_de_settings_y_no_del_codigo(self):
        """
        El nombre de la institución tiene que poder corregirse desde el
        .env: un reporte ejecutivo con el nombre mal impreso es peor que
        uno sin membrete.
        """
        with override_settings(
            ORG_NOMBRE="Otra Institución", ORG_UNIDAD="", ORG_CLASIFICACION="Reservado"
        ):
            self.assertEqual(membrete.encabezado_completo(), "Otra Institución")
            self.assertEqual(membrete.clasificacion(), "Reservado")

    def test_la_unidad_se_concatena_solo_si_esta_cargada(self):
        with override_settings(ORG_NOMBRE="Banco X", ORG_UNIDAD=""):
            self.assertEqual(membrete.encabezado_completo(), "Banco X")
        with override_settings(ORG_NOMBRE="Banco X", ORG_UNIDAD="Gerencia Y"):
            self.assertEqual(membrete.encabezado_completo(), "Banco X · Gerencia Y")

    def test_el_color_se_normaliza_para_openpyxl_y_pptx(self):
        with override_settings(ORG_COLOR="#203078"):
            self.assertEqual(membrete.color(), "203078")
            self.assertEqual(membrete.color_rgb(), (32, 48, 120))

    def test_el_logo_del_proyecto_existe(self):
        self.assertIsNotNone(
            membrete.ruta_logo(), "Falta app/tracker/marca/bdt-logo.png"
        )

    def test_sin_logo_no_revienta(self):
        """Un despliegue sin el archivo saca el reporte igual, sin logo."""
        with override_settings(ORG_LOGO="/no/existe/logo.png"):
            self.assertIsNone(membrete.ruta_logo())
            contenido = exportar_tickets(Ticket.objects.none())
            self.assertGreater(len(contenido), 0)


class DescribirFiltrosTest(TestCase):
    def test_sin_filtros_lo_dice_explicitamente(self):
        self.assertIn("Sin filtros", membrete.describir_filtros({}))
        self.assertIn(
            "Sin filtros",
            membrete.describir_filtros({"categoria": "", "q": ""}),
        )

    def test_traduce_los_valores_a_texto_legible(self):
        texto = membrete.describir_filtros(
            {"categoria": "proyecto", "riesgo": "critical"},
            categorias=Ticket.CATEGORIAS,
            niveles_riesgo=[("critical", "Crítico")],
        )
        self.assertIn("Tipo: Proyecto", texto)
        self.assertIn("Riesgo: Crítico", texto)


class MembreteEnElExcelTest(TestCase):
    def setUp(self):
        Ticket.objects.create(
            codigo="PRY-001", categoria=Ticket.PROYECTO, nombre="Proyecto",
            avance=40, fecha_fin=date(2026, 12, 31),
        )

    def _hoja(self, contenido, nombre_hoja="Proyectos"):
        ruta = a_disco(contenido)
        try:
            return load_workbook(ruta)[nombre_hoja]
        finally:
            os.unlink(ruta)

    def test_el_nombre_de_la_institucion_esta_en_la_hoja(self):
        ws = self._hoja(exportar_tickets(Ticket.objects.all()))
        textos = [
            c.value
            for fila in ws.iter_rows(min_row=1, max_row=FILA_ENCABEZADOS)
            for c in fila
            if isinstance(c.value, str)
        ]
        self.assertTrue(
            any(membrete.nombre() in t for t in textos),
            f"El membrete no aparece. Textos hallados: {textos}",
        )

    def test_la_clasificacion_esta_en_la_hoja(self):
        ws = self._hoja(exportar_tickets(Ticket.objects.all()))
        textos = " ".join(
            str(c.value)
            for fila in ws.iter_rows(min_row=1, max_row=FILA_ENCABEZADOS)
            for c in fila
            if c.value
        )
        self.assertIn(membrete.clasificacion(), textos)

    def test_registra_quien_exporto_y_con_que_filtros(self):
        usuario = get_user_model().objects.create_user(
            username="ana", password="x-larga-123"
        )
        contenido = exportar_tickets(
            Ticket.objects.all(),
            contexto={
                "usuario": usuario,
                "filtros_args": {
                    "filtros": {"categoria": "proyecto"},
                    "categorias": Ticket.CATEGORIAS,
                },
            },
        )
        ws = self._hoja(contenido)
        linea = str(ws.cell(row=5, column=1).value)

        self.assertIn("Emitido el", linea)
        self.assertIn("ana", linea)
        self.assertIn("Tipo: Proyecto", linea)

    def test_el_logo_queda_embebido(self):
        ruta = a_disco(exportar_tickets(Ticket.objects.all()))
        try:
            ws = load_workbook(ruta)["Proyectos"]
            self.assertEqual(len(ws._images), 1, "El logo no se embebió")
        finally:
            os.unlink(ruta)

    def test_las_propiedades_del_archivo_llevan_la_institucion(self):
        ruta = a_disco(exportar_tickets(Ticket.objects.all()))
        try:
            props = load_workbook(ruta).properties
            self.assertEqual(props.creator, membrete.encabezado_completo())
            self.assertEqual(props.description, membrete.clasificacion())
        finally:
            os.unlink(ruta)


class ElMembreteNoRompeElRoundTripTest(TestCase):
    """
    El membrete empuja los datos cuatro filas hacia abajo. Estos tests
    existen para que eso no se lleve puesto la reimportación.
    """

    def setUp(self):
        for codigo, cat, nombre in (
            ("PRY-001", Ticket.PROYECTO, "Proyecto uno"),
            ("REQ-001", Ticket.REQUERIMIENTO, "Requerimiento uno"),
            ("INC-001", Ticket.INCIDENCIA, "Incidencia una"),
        ):
            t = Ticket.objects.create(
                codigo=codigo, categoria=cat, nombre=nombre,
                avance=40, fecha_fin=date(2026, 12, 31),
            )
        Tarea.objects.create(
            ticket=Ticket.objects.get(codigo="PRY-001"),
            codigo="PRY-001.T01", descripcion="Analizar",
        )

    def test_los_datos_arrancan_donde_dice_la_constante(self):
        ruta = a_disco(exportar_tickets(Ticket.objects.all()))
        try:
            ws = load_workbook(ruta)["Proyectos"]
            self.assertEqual(ws.cell(row=FILA_DATOS, column=1).value, "PRY-001")
        finally:
            os.unlink(ruta)

    def test_el_importador_ignora_las_filas_del_membrete(self):
        ruta = a_disco(exportar_tickets(Ticket.objects.all()))
        try:
            filas, avisos = leer_filas(ruta)
        finally:
            os.unlink(ruta)

        self.assertEqual(len(filas), 3, f"avisos: {avisos}")
        self.assertEqual(avisos, [], "El membrete generó filas descartadas")
        self.assertEqual(
            {f["codigo"] for f in filas}, {"PRY-001", "REQ-001", "INC-001"}
        )

    def test_reimportar_sigue_sin_duplicar(self):
        ruta = a_disco(exportar_tickets(Ticket.objects.all()))
        try:
            batch = importar(ruta, "registros.xlsx")
        finally:
            os.unlink(ruta)

        self.assertEqual(Ticket.objects.count(), 3)
        self.assertEqual(batch.creados, 0)
        self.assertEqual(batch.filas_leidas, 3)

    def test_el_archivo_de_tareas_se_sigue_reconociendo(self):
        """
        Con el membrete los encabezados bajaron a la fila 6. La detección
        miraba solo hasta la 6: margen cero.
        """
        ruta = a_disco(exportar_tareas(Tarea.objects.all()))
        try:
            self.assertTrue(es_archivo_de_tareas(ruta))
            batch = importar_tareas(ruta, "tareas.xlsx")
        finally:
            os.unlink(ruta)

        self.assertEqual(Tarea.objects.count(), 1)
        self.assertEqual(batch.creados, 0)


class MembreteEnLaPresentacionTest(TestCase):
    def test_el_pptx_lleva_el_logo_en_todas_las_laminas(self):
        from tracker.presentacion import generar_presentacion_pptx
        from pptx import Presentation

        Ticket.objects.create(
            codigo="PRY-001", categoria=Ticket.PROYECTO, nombre="P", avance=50
        )
        contenido = generar_presentacion_pptx(
            periodo="mensual", tickets=Ticket.objects.all(), metricas=None
        )

        ruta = a_disco(contenido, sufijo=".pptx")
        try:
            prs = Presentation(ruta)
            con_imagen = [
                i
                for i, s in enumerate(prs.slides, start=1)
                if any(sh.shape_type == 13 for sh in s.shapes)
            ]
        finally:
            os.unlink(ruta)

        self.assertEqual(
            len(con_imagen), 7, f"Láminas con logo: {con_imagen} (se esperaban 7)"
        )


class ExportacionesDesdeLaVistaTest(TestCase):
    def setUp(self):
        self.user = get_user_model().objects.create_user(
            username="ana", password="clave-larga-123"
        )
        self.client.force_login(self.user)
        Ticket.objects.create(
            codigo="PRY-001", categoria=Ticket.PROYECTO, nombre="P", avance=50
        )

    def test_las_tres_exportaciones_llevan_membrete(self):
        urls = [
            reverse("exportar_registros"),
            reverse("exportar_tareas"),
            reverse("exportar_tareas_de", args=[Ticket.objects.get().pk]),
        ]

        for url in urls:
            with self.subTest(url=url):
                respuesta = self.client.get(url)
                self.assertEqual(respuesta.status_code, 200)

                ruta = a_disco(respuesta.content)
                try:
                    wb = load_workbook(ruta)
                    ws = wb[wb.sheetnames[0]]
                    textos = " ".join(
                        str(c.value)
                        for fila in ws.iter_rows(min_row=1, max_row=FILA_ENCABEZADOS)
                        for c in fila
                        if c.value
                    )
                    self.assertIn(membrete.nombre(), textos)
                    self.assertIn("ana", textos)
                finally:
                    os.unlink(ruta)
